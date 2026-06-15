"""Fill the MANDATORY Redrob 'Idea Submission' template with our project content,
preserving its branding. Output: Redrob_Idea_Submission.pptx"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC = "_template_redrob.pptx"
OUT = "Redrob_Idea_Submission.pptx"

DARK = RGBColor(0x20, 0x27, 0x29)
PURPLE = RGBColor(0x6B, 0x46, 0xE6)
BLUE = RGBColor(0x2E, 0x4B, 0xF0)
GREY = RGBColor(0x5A, 0x63, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LAV = RGBColor(0xF1, 0xEE, 0xFD)
LAVLN = RGBColor(0xD9, 0xD2, 0xFA)
GREEN = RGBColor(0x12, 0x9E, 0x6B)
RED = RGBColor(0xD8, 0x4A, 0x32)
BODY_FONT = "Manrope"

# ---------- placeholders the user must confirm ----------
TEAM_NAME = "Team-Phoenix"
TEAM_LEADER = "Akash Boddula"
GITHUB = "github.com/akashboddula2425-hub/AI-candidate-ranking-system"
SANDBOX = "Run locally — streamlit run app.py  (hosted Space optional; see README)"

prs = Presentation(SRC)


def body_shape(slide, title_text):
    """Return (title_shape, body_shape). Body = the largest text box below the title."""
    title_sh = None
    body = None
    best_h = -1
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        top_in = sh.top / 914400
        if 0.7 <= top_in <= 1.1 and title_sh is None:
            title_sh = sh
        elif top_in >= 1.25:
            if sh.height > best_h:
                best_h, body = sh.height, sh
    return title_sh, body


def set_bullets(tf, items, size=13, gap=8, lead_color=PURPLE):
    """items: list of (lead, rest). Renders '▸ lead  rest' paragraphs.
    Accepts a shape or a text_frame."""
    if hasattr(tf, "text_frame"):
        tf = tf.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, (lead, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
        # suppress any inherited bullet glyph so only our '▸' marker shows
        pPr = p._p.get_or_add_pPr()
        for tag in ("a:buChar", "a:buAutoNum", "a:buNone"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))
        r0 = p.add_run(); r0.text = "▸  "
        r0.font.size = Pt(size); r0.font.bold = True; r0.font.color.rgb = lead_color; r0.font.name = BODY_FONT
        if lead:
            r1 = p.add_run(); r1.text = lead
            r1.font.size = Pt(size); r1.font.bold = True; r1.font.color.rgb = DARK; r1.font.name = BODY_FONT
            if rest:
                r1.text = lead + "  "
        if rest:
            r2 = p.add_run(); r2.text = rest
            r2.font.size = Pt(size); r2.font.bold = False; r2.font.color.rgb = DARK; r2.font.name = BODY_FONT


def add_box(slide, x, y, w, h, fill, line=None, lw=1.0, rounded=True):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    try:
        shp.adjustments[0] = 0.12
    except Exception:
        pass
    return shp


def box_text(slide, x, y, w, h, lines, size=10, color=DARK, bold0=True, align=PP_ALIGN.CENTER):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Emu(0))
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = 1.0
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size if i else size + (1 if bold0 else 0))
        r.font.bold = (i == 0 and bold0); r.font.color.rgb = color; r.font.name = BODY_FONT
    return tb


def arrow(slide, x, y, w, color=PURPLE):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.18))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background(); a.shadow.inherit = False
    return a


S = prs.slides

# ===================== SLIDE 1 — title fields =====================
for sh in S[0].shapes:
    if sh.has_text_frame:
        t = sh.text_frame.text.strip()
        val = None
        if t.startswith("Team Name"):
            val = TEAM_NAME
        elif t.startswith("Team Leader"):
            val = TEAM_LEADER
        elif t.startswith("Problem Statement"):
            val = ("Intelligent Candidate Discovery & Ranking — rank the best-fit "
                   "candidates for a job description by genuine fit, not keywords.")
        if val is not None:
            p = sh.text_frame.paragraphs[0]
            base = p.runs[0]
            base.font.color.rgb = DARK
            r = p.add_run(); r.text = "  " + val
            r.font.bold = True; r.font.color.rgb = PURPLE
            r.font.name = "Manrope SemiBold"
            if base.font.size:
                r.font.size = base.font.size

# ===================== SLIDE 2 — Solution Overview =====================
_, b = body_shape(S[1], "Solution Overview")
set_bullets(b, [
    ("A hybrid candidate ranker.", "It reads each candidate's full profile and ranks the top 100 by genuine fit — combining AI semantic understanding with a structured, recruiter-style scoring model."),
    ("Four layers working together:", "(1) dense embeddings that grasp meaning, (2) structured fit rules, (3) integrity guards that reject fakes, (4) behavioural signals for real availability."),
    ("What makes it different:", "AI-skill credit is GATED by title/career credibility — so a Marketing Manager who lists 'RAG' earns almost nothing, while a genuine engineer earns full weight."),
    ("Sees plain-language fits:", "surfaces people who built ranking/retrieval systems but never wrote the buzzword — which keyword filters miss entirely."),
    ("Trustworthy & fast:", "every pick has an explainable score + honest reason; runs in ~37s on a CPU with no paid LLM API calls."),
], size=12.5)

# ===================== SLIDE 3 — JD Understanding & Candidate Evaluation =====================
_, b = body_shape(S[2], "JD Understanding")
set_bullets(b, [
    ("Requirements extracted from the JD:", "production embeddings & vector search; a shipped ranking / recsys / search system at scale; product-company (not services-only); 5–9 yrs; evaluation rigor (NDCG/MRR/MAP); Noida/Pune; a pragmatic shipper, not research-only."),
    ("Most important candidate signals:", "current + past job titles, what they actually built (career-history text), real retrieval/ranking evidence, experience band, product-vs-services history, and behavioural availability."),
    ("Fit beyond keywords:", "we embed a distilled version of the JD and match each profile by meaning; we read career descriptions, not just the skills list; and we gate buzzword credit by whether the title/career makes it believable."),
], size=12.5)

# ===================== SLIDE 4 — Ranking Methodology =====================
_, b = body_shape(S[3], "Ranking Methodology")
set_bullets(b, [
    ("Retrieve:", "every profile is pre-embedded with a sentence-transformer (all-MiniLM-L6-v2, 384-d) and indexed with FAISS; we score cosine similarity to a distilled JD query."),
    ("Score (weighted blend 'fit'):", "title .22 · what-they-built .22 · career-arc .14 · semantic-retrieval .13 · semantic-full .11 · experience .10 · location .08."),
    ("Combine all signals:", "final = fit × behavioural × disqualifier × honeypot × stuffer — multipliers that sink unavailable, off-domain, fake, or impossible profiles."),
    ("Rank:", "sort by final score (deterministic tie-break by candidate_id) and emit the top 100 with reasons."),
], size=12.5)

# ===================== SLIDE 5 — Explainability & Data Validation =====================
_, b = body_shape(S[4], "Explainability")
set_bullets(b, [
    ("Every decision is explained:", "each shortlisted candidate gets a 1–2 sentence reason built from their own facts — title, years, real employers, named skills — plus honest concerns (e.g. long notice, low response rate)."),
    ("No hallucinations:", "reasoning is composed only from fields actually present in the profile (no LLM, no invented skills); tone is tied to the rank so it never contradicts the score."),
    ("Suspicious / low-quality profiles:", "honeypots (impossible tenure vs dates; 'expert' in many skills with 0 months used) are detected by simple arithmetic and crushed ×0.03; keyword-stuffers ×0.25–0.40."),
    ("Reproducible:", "output is fully deterministic — two runs produce a byte-identical file."),
], size=12.5)

# ===================== SLIDE 6 — End-to-End Workflow (diagram) =====================
_, b = body_shape(S[5], "End-to-End Workflow")
set_bullets(b, [("From JD to ranked candidates:", "one offline preparation step, then a fast CPU-only ranking pass.")], size=12.5)
# linear flow of step boxes
steps = ["Job\nDescription", "Distil to query\n+ lexicons", "Embed 100K\nprofiles (offline)",
         "Score: semantic +\nstructured + guards", "Combine ×\nmultipliers", "Sort →\nTop 100", "submission.csv\n+ reasons"]
cols = [PURPLE, LAV, LAV, LAV, LAV, LAV, GREEN]
x = 0.34; y = 2.5; bw = 1.24; gap = 0.10; bh = 1.0
for i, (txt, c) in enumerate(zip(steps, cols)):
    dark_fill = c in (PURPLE, GREEN)
    add_box(S[5], x, y, bw, bh, c, None if dark_fill else LAVLN)
    box_text(S[5], x + 0.06, y, bw - 0.12, bh, txt.split("\n"),
             size=9, color=(WHITE if dark_fill else DARK))
    if i < len(steps) - 1:
        arrow(S[5], x + bw + 0.005, y + bh / 2 - 0.09, gap - 0.01)
    x += bw + gap

# ===================== SLIDE 7 — System Architecture (diagram) =====================
# OFFLINE band
add_box(S[6], 0.34, 1.55, 2.0, 2.95, RGBColor(0x10, 0x26, 0x2A))
box_text(S[6], 0.42, 1.62, 1.84, 0.6, ["OFFLINE (once)", "build_index.py"], size=11, color=RGBColor(0x00, 0xC2, 0xA8))
box_text(S[6], 0.42, 2.3, 1.84, 2.1,
         ["Read every profile", "↓", "all-MiniLM-L6-v2", "↓", "embeddings.f16.npy", "+ FAISS index"],
         size=9.5, color=WHITE, bold0=False)
arrow(S[6], 2.4, 2.9, 0.22)
# RANKING band with 4 lanes
add_box(S[6], 2.74, 1.55, 6.92, 2.95, RGBColor(0xF7, 0xF7, 0xFB), LAVLN)
box_text(S[6], 2.9, 1.62, 6.0, 0.4, ["RANKING  ·  rank.py  ·  ~37 s  ·  CPU only, no network"], size=11, color=DARK)
lanes = [("Dense semantic", "cosine to JD query", PURPLE),
         ("Structured fit", "title · built · exp · loc", BLUE),
         ("Integrity guards", "honeypot & stuffer kill", RED),
         ("Behavioural", "availability multiplier", GREEN)]
lx = 2.92; lw = 1.62
for i, (t, d, c) in enumerate(lanes):
    add_box(S[6], lx, 2.2, lw, 1.0, c)
    box_text(S[6], lx + 0.06, 2.26, lw - 0.12, 1.0, [t, d], size=9.5, color=WHITE)
    lx += lw + 0.08
# combine + output
add_box(S[6], 2.92, 3.45, 3.3, 0.85, RGBColor(0x1E, 0x27, 0x61))
box_text(S[6], 2.98, 3.45, 3.2, 0.85, ["Combine → final score", "fit × behavioural × guards"], size=10, color=WHITE)
arrow(S[6], 6.3, 3.78, 0.22)
add_box(S[6], 6.6, 3.45, 3.0, 0.85, GREEN)
box_text(S[6], 6.66, 3.45, 2.9, 0.85, ["Top 100 + reasoning", "→ submission.csv"], size=10, color=WHITE)

# ===================== SLIDE 8 — Results & Performance =====================
_, b = body_shape(S[7], "Results & Performance")
set_bullets(b, [
    ("Ranking quality (released 100K pool):", "top-100 is 100% engineering titles; 86/100 sit in the ideal 5–9-yr band (mean 6.5); Noida & Pune are the #1 and #2 locations — exactly the JD's preference."),
    ("Trap resistance:", "0 honeypots and 0 keyword-stuffers in the top 100 (the spec disqualifies any submission with >10% honeypots)."),
    ("Meets every compute constraint:", "~37 s for 100K candidates on a 12-core CPU (limit is 5 min), ≤16 GB RAM, CPU-only, no network and no LLM API calls during ranking."),
    ("Production-shaped:", "heavy embedding is a one-time offline step (~20 min, allowed); the ranking step is pure fast array math and is byte-for-byte reproducible."),
], size=12.5)

# ===================== SLIDE 9 — Technologies Used =====================
_, b = body_shape(S[8], "Technologies Used")
set_bullets(b, [
    ("Python 3", "— the whole system; clean and auditable."),
    ("sentence-transformers (all-MiniLM-L6-v2) + PyTorch (CPU)", "— turns each profile into a meaning-vector so we match by intent, not keywords; small and CPU-friendly."),
    ("FAISS", "— Facebook's similarity search; the vector-database / retrieval path."),
    ("NumPy", "— all scoring is fast array math, which is why ranking takes seconds."),
    ("scikit-learn", "— TF-IDF fallback so the pipeline runs even without the model present."),
    ("Streamlit", "— the interactive sandbox demo. Why this stack: free, lightweight, CPU-only, no paid APIs, fully reproducible."),
], size=11.5, gap=6)

# ===================== SLIDE 10 — Submission Assets =====================
_, b = body_shape(S[9], "Submission Assets")
set_bullets(b, [
    ("GitHub repository:", GITHUB + "  (code, README, requirements, reproduce command)."),
    ("Sandbox / demo:", SANDBOX),
    ("Ranked output:", "submission.csv  (top 100 candidates, validated format)."),
    ("Reproduce command:", "python rank.py --candidates ./candidates.jsonl --out ./submission.csv"),
    ("Documentation:", "approach deck + detailed project guide (PDF) included in /docs."),
], size=12.5)

prs.save(OUT)
print("saved", OUT, "with", len(prs.slides._sldIdLst), "slides")
