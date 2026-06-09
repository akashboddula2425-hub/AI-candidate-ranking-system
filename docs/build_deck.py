"""Redrob approach deck -> approach_deck.pptx (python-pptx; opens cleanly in PowerPoint)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# palette
INK = RGBColor(0x0F, 0x1B, 0x3C)
NAVY = RGBColor(0x1E, 0x27, 0x61)
NAVY2 = RGBColor(0x16, 0x22, 0x4A)
NAVY3 = RGBColor(0x26, 0x35, 0x6E)
TEAL = RGBColor(0x00, 0xC2, 0xA8)
CYAN = RGBColor(0x3D, 0xA9, 0xFC)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
SLATE = RGBColor(0x33, 0x41, 0x5C)
MUTED = RGBColor(0x7C, 0x87, 0x9B)
CARD = RGBColor(0xF4, 0xF7, 0xFB)
CARDLINE = RGBColor(0xE2, 0xE8, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xF4, 0x79, 0x5B)
AMBER = RGBColor(0xF4, 0xA9, 0x3B)
VIOLET = RGBColor(0x9B, 0x8C, 0xFF)
DARKRED = RGBColor(0x2A, 0x1A, 0x22)
DARKREDLN = RGBColor(0x5E, 0x2B, 0x33)
DARKTEAL = RGBColor(0x10, 0x26, 0x2A)
DARKTEALLN = RGBColor(0x1F, 0x5E, 0x54)

HF = "Georgia"
BF = "Calibri"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def _noline(shp):
    shp.line.fill.background()


def _line(shp, color, w=1.0):
    shp.line.color.rgb = color
    shp.line.width = Pt(w)


def rect(s, x, y, w, h, fill, line=None, lw=1.0, rounded=True, radius=0.08):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        _noline(shp)
    else:
        _line(shp, line, lw)
    shp.shadow.inherit = False
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    return shp


def oval(s, x, y, w, h, fill):
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    _noline(shp)
    shp.shadow.inherit = False
    return shp


def text(s, x, y, w, h, runs, size=14, bold=False, color=SLATE, font=BF,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, italic=False, spacing=1.0):
    """runs: a string (may contain \n) -> paragraphs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Emu(0))
    lines = runs.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = ln
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.name = font
        f.color.rgb = color
    return tb


def head(s, kicker, title, color=INK):
    text(s, 0.7, 0.5, 11.9, 0.3, kicker.upper(), size=13, bold=True, color=TEAL, font=BF)
    text(s, 0.7, 0.82, 11.9, 1.0, title, size=29, bold=True, color=color, font=HF, spacing=0.98)


def chip(s, x, y, n, color=TEAL):
    oval(s, x, y, 0.5, 0.5, color)
    # contrast the number against the chip fill
    num_color = WHITE if color in (INK, NAVY, NAVY2) else INK
    text(s, x, y, 0.5, 0.5, str(n), size=18, bold=True, color=num_color, font=HF,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


# ===================================================== 1 TITLE
s = add_slide(INK)
oval(s, 9.7, -2.2, 6.5, 6.5, NAVY)
oval(s, 11.2, 3.9, 4.2, 4.2, RGBColor(0x16, 0x20, 0x4A))
text(s, 0.8, 1.45, 11, 0.35, "REDROB · INTELLIGENT CANDIDATE DISCOVERY & RANKING",
     size=14, bold=True, color=TEAL)
text(s, 0.8, 2.05, 11.6, 1.2, "Reading between the lines", size=50, bold=True, color=WHITE, font=HF)
text(s, 0.82, 3.5, 10.7, 1.0,
     "A candidate ranker that hires the way a great recruiter does —\n"
     "by understanding what the role means, not by matching keywords.",
     size=18, color=ICE, spacing=1.15)
stats = [("100K", "candidates"), ("Top 100", "trusted shortlist"),
         ("Hybrid", "dense + reasoning"), ("36 s", "CPU, no network")]
for i, (a, b) in enumerate(stats):
    x = 0.85 + i * 3.0
    text(s, x, 5.35, 2.8, 0.6, a, size=30, bold=True, color=TEAL, font=HF)
    text(s, x, 5.98, 2.8, 0.4, b, size=13, color=ICE)

# ===================================================== 2 PROBLEM
s = add_slide(WHITE)
head(s, "The problem", "Keyword filters can't see what actually matters")
rows = [
    ("Recruiters scan hundreds of profiles",
     "and still miss the right person — not for lack of talent, but because filters match words, not fit."),
    ("The JD says one thing and means another",
     "“5–9 years” is a guide, not a gate. “AI skills” without the right title is a red flag, not a green one."),
    ("The best fit may never say the buzzword",
     "Someone who built “the ranking systems that decide what users see” is a fit even if they never write “RAG”."),
]
for i, (h, b) in enumerate(rows):
    y = 2.2 + i * 1.5
    chip(s, 0.75, y + 0.05, i + 1)
    text(s, 1.5, y, 11.0, 0.5, h, size=19, bold=True, color=INK)
    text(s, 1.5, y + 0.5, 10.9, 0.8, b, size=15.5, color=SLATE, spacing=1.05)

# ===================================================== 3 DATA & GAP
s = add_slide(WHITE)
head(s, "The brief", "One narrow role, a 100,000-profile haystack")
rect(s, 0.7, 2.0, 6.0, 4.7, CARD, CARDLINE)
text(s, 1.0, 2.22, 5.4, 0.4, "What the JD really wants", size=16, bold=True, color=TEAL)
want = [
    "Production embeddings retrieval + vector search",
    "Shipped a ranking / recsys / search system at scale",
    "At a product company (not services-only)",
    "Evaluation rigor: NDCG, MRR, MAP, A/B testing",
    "5–9 yrs, pragmatic shipper > research-only",
    "In or willing to relocate to Noida / Pune, India",
]
for i, t in enumerate(want):
    text(s, 1.0, 2.78 + i * 0.6, 0.3, 0.4, "›", size=16, bold=True, color=TEAL, font=HF)
    text(s, 1.34, 2.8 + i * 0.6, 5.1, 0.55, t, size=14.5, color=SLATE)
cells = [("100,000", "candidate profiles", TEAL), ("~1,000", "genuine ML/AI engineers", CYAN),
         ("~5,500", "keyword-stuffer traps", RED), ("~80", "impossible honeypots", RED)]
for i, (a, b, c) in enumerate(cells):
    x = 7.0 + (i % 2) * 3.0
    y = 2.0 + (i // 2) * 1.55
    rect(s, x, y, 2.8, 1.4, CARD, CARDLINE)
    text(s, x + 0.05, y + 0.16, 2.7, 0.7, a, size=29, bold=True, color=c, font=HF, align=PP_ALIGN.CENTER)
    text(s, x + 0.1, y + 0.9, 2.6, 0.45, b, size=12, color=SLATE, align=PP_ALIGN.CENTER)
rect(s, 7.0, 5.15, 5.8, 1.55, INK)
text(s, 7.25, 5.3, 5.4, 0.4, "Scored on a hidden, graded ground truth", size=14, bold=True, color=TEAL)
text(s, 7.25, 5.74, 5.45, 0.5, "0.50·NDCG@10 + 0.30·NDCG@50 + 0.15·MAP + 0.05·P@10",
     size=12.5, color=WHITE, font=MONO)
text(s, 7.25, 6.2, 5.4, 0.4, "Top-10 quality is half the score — precision at the very top is everything.",
     size=11.5, italic=True, color=ICE)

# ===================================================== 4 TRAPS CONTRAST
s = add_slide(INK)
head(s, "The decisive contrast", "Same keywords. Opposite verdicts.", WHITE)
rect(s, 0.7, 2.05, 5.9, 4.6, DARKRED, DARKREDLN)
text(s, 1.0, 2.22, 5.3, 0.4, "✕  KEYWORD-STUFFER", size=15, bold=True, color=RED)
text(s, 1.0, 2.68, 5.4, 0.6, "“Project Manager · AI enthusiast · Building with LLMs”",
     size=13.5, italic=True, color=ICE)
for i, t in enumerate(["Title & career: PM / Marketing / Sales",
                       "Skills list sprayed with RAG, Pinecone, FAISS",
                       "Summary: “took online courses on RAG…”",
                       "→ buzzwords are cheap; the work isn't there"]):
    text(s, 1.0, 3.45 + i * 0.62, 5.4, 0.55, t, size=14, color=WHITE)
text(s, 1.0, 6.05, 5.4, 0.4, "our score: 0.04  ·  buried", size=14, bold=True, color=RED, font=MONO)
rect(s, 6.75, 2.05, 5.9, 4.6, DARKTEAL, DARKTEALLN)
text(s, 7.05, 2.22, 5.3, 0.4, "✓  PLAIN-LANGUAGE TIER-5", size=15, bold=True, color=TEAL)
text(s, 7.05, 2.68, 5.4, 0.6, "“…the ranking and retrieval systems that decide what to show”",
     size=13.5, italic=True, color=ICE)
for i, t in enumerate(["Title & career: Senior AI Eng @ product co.",
                       "Skills in plain words: Ranking Systems, BM25",
                       "Summary maps 1:1 onto the JD's mandate",
                       "→ never says “RAG”, but clearly built it"]):
    text(s, 7.05, 3.45 + i * 0.62, 5.4, 0.55, t, size=14, color=WHITE)
text(s, 7.05, 6.05, 5.4, 0.4, "our score: 0.80  ·  top of the list", size=14, bold=True, color=TEAL, font=MONO)

# ===================================================== 5 KEY INSIGHT
s = add_slide(WHITE)
head(s, "The insight that drives everything", "Buzzwords are noise. Evidence and availability are signal.")
cards = [
    ("Read the career, not the skills list",
     "Credit for an AI skill is GATED by title/career credibility. A Marketing Manager's “RAG” earns almost nothing; an engineer's earns full weight."),
    ("A quiet tell in the data",
     "“Python” appears as a listed skill on only ~1.4% of profiles — the genuine engineers. Stuffers reach for the flashy buzzwords instead."),
    ("Hireable ≠ qualified",
     "A perfect profile that hasn't logged in for 6 months with a 5% response rate isn't available. Behaviour modifies fit — it never manufactures it."),
]
glyph = ["⤷", "◔", "◉"]
for i, (h, b) in enumerate(cards):
    x = 0.7 + i * 4.05
    rect(s, x, 2.2, 3.75, 4.3, CARD, CARDLINE)
    rect(s, x + 0.3, 2.5, 0.85, 0.85, INK, radius=0.18)
    text(s, x + 0.3, 2.5, 0.85, 0.85, glyph[i], size=24, color=TEAL,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.3, 3.55, 3.2, 0.85, h, size=17, bold=True, color=INK, spacing=0.98)
    text(s, x + 0.3, 4.45, 3.2, 1.9, b, size=14, color=SLATE, spacing=1.08)

# ===================================================== 6 ARCHITECTURE
s = add_slide(INK)
head(s, "Architecture", "A hybrid: dense recall, structured precision", WHITE)
stages = [
    ("Profile", "headline, summary,\ncareer, skills, signals", NAVY),
    ("Dense semantic", "all-MiniLM-L6-v2 + FAISS\ncosine to distilled JD", TEAL),
    ("Structured fit", "title · what they built ·\nexperience · career · location", CYAN),
    ("Guards + behaviour", "honeypot & stuffer kill ×\navailability multiplier", AMBER),
    ("Ranked top 100", "score + honest\nreasoning", TEAL),
]
cw, gap, x0, y = 2.3, 0.28, 0.62, 2.45
for i, (t, d, c) in enumerate(stages):
    x = x0 + i * (cw + gap)
    rect(s, x, y, cw, 2.5, NAVY2, NAVY3)
    rect(s, x + 0.18, y + 0.22, 0.55, 0.16, c, radius=0.3)
    text(s, x + 0.18, y + 0.48, cw - 0.36, 0.7, t, size=15, bold=True, color=WHITE, spacing=0.95)
    text(s, x + 0.18, y + 1.22, cw - 0.36, 1.1, d, size=11, color=ICE, spacing=1.05)
    if i < len(stages) - 1:
        text(s, x + cw - 0.06, y + 0.95, 0.34, 0.5, "→", size=22, bold=True, color=TEAL,
             font=HF, align=PP_ALIGN.CENTER)
rect(s, 0.62, 5.35, 12.1, 1.35, DARKTEAL, DARKTEALLN)
text(s, 0.9, 5.5, 2.2, 0.4, "Why hybrid?", size=15, bold=True, color=TEAL)
text(s, 0.9, 5.86, 11.6, 0.75,
     "Embeddings give recall over fits described in plain language. The structured + integrity layers give precision and "
     "keep stuffers and honeypots out of the top. Neither wins alone — which is exactly the system this role is hired to build.",
     size=13.5, color=WHITE, spacing=1.05)

# ===================================================== 7 SCORING
s = add_slide(WHITE)
head(s, "The scoring model", "One transparent, auditable formula")
rect(s, 0.7, 2.0, 11.95, 0.95, INK)
text(s, 0.7, 2.0, 11.95, 0.95, "final  =  fit  ×  behavioural  ×  disqualifier  ×  honeypot  ×  stuffer",
     size=18, bold=True, color=WHITE, font=MONO, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
text(s, 0.7, 3.22, 6, 0.4, "fit  =  weighted blend", size=16, bold=True, color=INK)
weights = [("Title (current + career)", 0.22), ("What they built (lexical)", 0.22),
           ("Career arc · product vs services", 0.14), ("Semantic · retrieval facet", 0.13),
           ("Semantic · full JD", 0.11), ("Experience band (5–9y)", 0.10), ("Location", 0.08)]
maxw, barX, barMax = 0.22, 3.7, 2.7
for i, (t, w) in enumerate(weights):
    y = 3.78 + i * 0.44
    text(s, 0.7, y, 3.0, 0.36, t, size=12.5, color=SLATE)
    rect(s, barX, y + 0.04, barMax, 0.24, CARD, radius=0.4)
    rect(s, barX, y + 0.04, barMax * (w / maxw), 0.24, TEAL, radius=0.4)
    text(s, barX + barMax + 0.1, y, 0.6, 0.36, f"{w:.2f}", size=11.5, bold=True, color=INK, font=MONO)
rect(s, 7.45, 3.22, 5.2, 3.5, CARD, CARDLINE)
text(s, 7.7, 3.4, 4.7, 0.4, "Multipliers", size=16, bold=True, color=INK)
mults = [("Behavioural availability", "× 0.62 – 1.08", TEAL),
         ("CV / speech, no IR-NLP", "× 0.55", RED),
         ("Research-only, no production", "× 0.60", RED),
         ("Keyword-stuffer", "× 0.25 – 0.40", RED),
         ("Honeypot (impossible profile)", "× 0.03", RED)]
for i, (t, v, c) in enumerate(mults):
    y = 3.95 + i * 0.55
    text(s, 7.7, y, 3.3, 0.45, t, size=13, color=SLATE)
    text(s, 10.85, y, 1.65, 0.45, v, size=12.5, bold=True, color=c, font=MONO, align=PP_ALIGN.RIGHT)

# ===================================================== 8 TRAP DEFENSE
s = add_slide(WHITE)
head(s, "Trap defense", "Every trap has a dedicated guard")
cols = [0.7, 3.4, 6.6, 12.65]
for i, t in enumerate(["TRAP", "HOW IT HIDES", "GUARD"]):
    text(s, cols[i], 1.95, cols[i + 1] - cols[i] - 0.2, 0.35, t, size=12, bold=True, color=MUTED)
data = [
    ("Keyword-stuffer", "Non-eng title + buzzword skills", "Gate skill-credit on title credibility; anti-query dampener", TEAL),
    ("Honeypot", "Tenure > dates allow; “expert”, 0 months used", "Date-arithmetic & mastery-without-time checks → ×0.03", RED),
    ("CV / speech specialist", "YOLO, OpenCV, ASR — no retrieval", "Penalise CV-heavy profiles lacking IR/NLP evidence", CYAN),
    ("Research-only", "Papers, labs, no deployment", "Require production signal or down-weight", AMBER),
    ("Title-chaser", "Job-hops every <18 months for titles", "Tenure-stability term in the career-arc score", VIOLET),
]
for i, row in enumerate(data):
    y = 2.45 + i * 0.86
    rect(s, 0.7, y, 11.95, 0.74, WHITE if i % 2 else CARD, CARDLINE)
    rect(s, 0.85, y + 0.22, 0.12, 0.3, row[3], radius=0.3)
    text(s, 1.1, y, 2.3, 0.74, row[0], size=14, bold=True, color=INK, valign=MSO_ANCHOR.MIDDLE)
    text(s, 3.4, y, 3.1, 0.74, row[1], size=12.5, color=SLATE, valign=MSO_ANCHOR.MIDDLE)
    text(s, 6.6, y, 5.9, 0.74, row[2], size=12.5, color=SLATE, valign=MSO_ANCHOR.MIDDLE)

# ===================================================== 9 REASONING
s = add_slide(WHITE)
head(s, "Trustworthy by design", "Every pick comes with honest, grounded reasoning")
rect(s, 0.7, 2.05, 7.0, 4.6, CARD, CARDLINE)
text(s, 1.0, 2.28, 6.4, 0.35, "GENERATED REASONING (rank 65)", size=12, bold=True, color=TEAL)
text(s, 1.0, 2.72, 6.4, 2.0,
     "“AI Engineer with 5 yrs (Salesforce); direct ranking systems, recommendation systems experience "
     "(Haystack, LoRA, Weaviate); Pune-based. Concern: long notice period (120d).”",
     size=16.5, color=INK, italic=True, spacing=1.15)
text(s, 1.0, 5.95, 6.4, 0.5, "Facts pulled only from the profile — no LLM, no invented skills.",
     size=13, color=MUTED)
checks = [
    "Specific facts: title, years, real employers",
    "Connects to JD: names the retrieval/ranking work",
    "Honest: flags the 120-day notice concern",
    "No hallucination: skills are from the profile",
    "Varied: composed per candidate, not templated",
    "Tone matches rank: confident high, hedged low",
]
for i, t in enumerate(checks):
    y = 2.15 + i * 0.74
    oval(s, 8.0, y + 0.03, 0.42, 0.42, TEAL)
    text(s, 8.0, y + 0.03, 0.42, 0.42, "✓", size=15, bold=True, color=INK,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(s, 8.6, y, 4.2, 0.5, t, size=13.5, color=SLATE, valign=MSO_ANCHOR.MIDDLE)

# ===================================================== 10 RESULTS
s = add_slide(INK)
head(s, "Results on the released pool", "A shortlist a recruiter can trust", WHITE)
big = [("100%", "engineering titles in the top 100", TEAL),
       ("86 / 100", "inside the 5–9 yr band (mean 6.5)", CYAN),
       ("0", "honeypots & 0 stuffers in top 100", TEAL),
       ("Noida + Pune", "the #1 and #2 locations", CYAN)]
for i, (a, b, c) in enumerate(big):
    x = 0.7 + (i % 2) * 6.05
    y = 2.25 + (i // 2) * 1.7
    rect(s, x, y, 5.8, 1.5, NAVY2, NAVY3)
    text(s, x + 0.35, y + 0.22, 5.2, 0.7, a, size=33, bold=True, color=c, font=HF)
    text(s, x + 0.37, y + 1.0, 5.2, 0.4, b, size=14, color=ICE)
rect(s, 0.7, 5.75, 11.95, 0.95, DARKTEAL, DARKTEALLN)
text(s, 0.95, 5.75, 11.5, 0.95,
     "Disqualification-proof: the spec fails any submission with >10% honeypots in the top 100. Ours has zero.",
     size=14.5, color=WHITE, valign=MSO_ANCHOR.MIDDLE)

# ===================================================== 11 COMPUTE
s = add_slide(WHITE)
head(s, "Built for production, not just a benchmark", "Fast, reproducible, no API in the loop")
left = [
    ("Offline · once", "Embed 100K profiles with a small CPU transformer; build the FAISS index. ~20 min, fully documented."),
    ("Ranking step · every run", "Pure numpy over the pre-computed index. 36 seconds for 100K on a 12-core CPU — far inside the 5-min budget."),
    ("No network, no GPU, no LLM calls", "The ranking code makes zero API calls — the same constraint a 200K-pool production system must meet."),
]
for i, (h, b) in enumerate(left):
    y = 2.2 + i * 1.5
    chip(s, 0.75, y + 0.05, i + 1, TEAL if i == 2 else INK)
    text(s, 1.5, y, 6.4, 0.5, h, size=17, bold=True, color=INK)
    text(s, 1.5, y + 0.5, 6.5, 0.9, b, size=14, color=SLATE, spacing=1.08)
rect(s, 8.4, 2.2, 4.25, 4.5, INK)
text(s, 8.7, 2.42, 3.7, 0.35, "ONE COMMAND", size=12, bold=True, color=TEAL)
text(s, 8.7, 2.88, 3.8, 1.3,
     "python rank.py\n  --candidates candidates.jsonl\n  --out submission.csv",
     size=12, color=WHITE, font=MONO, spacing=1.12)
specs = [("Runtime", "≤ 5 min  ✓ 36 s"), ("Memory", "≤ 16 GB  ✓"),
         ("Compute", "CPU only  ✓"), ("Network", "off  ✓")]
for i, (k, v) in enumerate(specs):
    y = 4.4 + i * 0.55
    text(s, 8.7, y, 1.7, 0.45, k, size=13, color=ICE)
    text(s, 10.2, y, 2.2, 0.45, v, size=12.5, bold=True, color=TEAL, font=MONO, align=PP_ALIGN.RIGHT)

# ===================================================== 12 CLOSING
s = add_slide(INK)
oval(s, -2.3, 3.6, 6.5, 6.5, NAVY)
text(s, 0.9, 1.3, 11, 0.35, "WHY IT WINS", size=14, bold=True, color=TEAL)
text(s, 0.9, 1.8, 11.5, 1.0, "Understanding, not keyword counting", size=37, bold=True, color=WHITE, font=HF)
pts = [
    "Hybrid dense + structured ranking — recall and precision, not one at the cost of the other.",
    "Engineered against every trap the dataset hides; zero honeypots and zero stuffers in the top 100.",
    "Optimised where it counts: top-10 quality is half the score, so the top is all in-band, on-domain fits.",
    "Fast, reproducible, and fully auditable — every score and every reason can be explained line by line.",
]
for i, t in enumerate(pts):
    y = 3.2 + i * 0.82
    oval(s, 0.95, y + 0.06, 0.32, 0.32, TEAL)
    text(s, 1.5, y, 11.0, 0.7, t, size=16.5, color=ICE, spacing=1.05)
text(s, 0.9, 6.72, 11.5, 0.5,
     "Make hiring smarter — rank candidates the way a great recruiter would.",
     size=15, italic=True, bold=True, color=WHITE)

prs.save("approach_deck.pptx")
print("wrote approach_deck.pptx with", len(prs.slides._sldIdLst), "slides")
